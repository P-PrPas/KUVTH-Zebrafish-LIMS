"""Interpret source cells without changing their immutable import draft values."""

from __future__ import annotations

import re
from datetime import date, datetime, timedelta
from typing import Any

from openpyxl.utils import get_column_letter


def _text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _day(value: Any) -> date | None:
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value)[:10])
    except ValueError:
        match = re.match(r"^(\d{4})[-/](\d{1,2})[-/](\d{1,2})", _text(value))
        if match:
            try:
                return date(*(int(part) for part in match.groups()))
            except ValueError:
                return None
        return None


def _clock(value: Any) -> str | None:
    raw = _text(value)
    match = re.match(r"^(\d{1,2})[.:](\d{2})(?:\D|$)", raw)
    if not match:
        return None
    hour, minute = (int(part) for part in match.groups())
    return f"{hour:02d}:{minute:02d}" if hour < 24 and minute < 60 else None


def _condition(normal: Any, abnormal: Any) -> str:
    if _text(normal) in {"1", "1.0"} and _text(abnormal) in {"0", "0.0", ""}:
        return "NORMAL"
    if _text(abnormal) in {"1", "1.0"} and _text(normal) in {"0", "0.0", ""}:
        return "ABNORMAL"
    return "UNDETERMINED"


def _flags(cells: dict[str, Any], start: int, end: int, dob: date | None,
           disposition: str, exit_date: date | None) -> tuple[list[dict[str, Any]], list[str]]:
    entries: list[dict[str, Any]] = []
    warnings: list[str] = []
    for index in range(start, end + 1):
        column = get_column_letter(index)
        if column not in cells:
            continue
        raw = cells[column]
        value = _text(raw)
        if value not in {"0", "1", "0.0", "1.0"}:
            warnings.append(f"{column}: expected 1 or 0")
            continue
        day_number = index - start + 1
        observed = dob + timedelta(days=day_number) if dob else None
        if value in {"1", "1.0"}:
            outcome = "ALIVE"
        elif disposition in {"FROZEN", "DISCARDED"} and exit_date and observed and observed >= exit_date:
            outcome = "TRACKING_ENDED"
        elif disposition in {"FROZEN", "DISCARDED"} and not exit_date:
            outcome = "UNRESOLVED_ZERO"
        else:
            outcome = "NOT_SURVIVING"
        entries.append({
            "day": day_number, "sourceColumn": column,
            "observedOn": observed.isoformat() if observed else None,
            "sourceValue": raw, "outcome": outcome,
        })
    return entries, warnings


def _fish(cells: dict[str, Any]) -> dict[str, Any]:
    v1 = bool(re.match(r"^(?:CLA|CL|RT|DC)\d+", _text(cells.get("B")), re.IGNORECASE)) or (
        bool(_text(cells.get("K"))) and _text(cells.get("K")).casefold() not in {"normal", "abnormal"}
    )
    dob = _day(cells.get("F" if v1 else "B"))
    exit_date = _day(cells.get("OD" if v1 else "J"))
    source_status = _text(cells.get("L")) if v1 else ("Frozen" if exit_date else "")
    source_status_folded = source_status.casefold()
    if v1:
        disposition = (
            "FROZEN" if source_status_folded == "frozen" else
            "DISCARDED" if source_status_folded == "discarded" else "NONE"
        )
        life_state = "DEAD" if source_status_folded == "dead" else "UNKNOWN"
    else:
        disposition = "FROZEN" if exit_date else "UNKNOWN"
        life_state = "UNKNOWN"
    status = (
        "DEAD" if life_state == "DEAD" and exit_date else
        disposition if disposition in {"FROZEN", "DISCARDED"} and exit_date else "UNKNOWN"
    )
    observations, warnings = _flags(cells, 19 if v1 else 14, 388 if v1 else 205,
                                    dob, disposition, exit_date)
    if dob is None:
        warnings.append("Date of birth is missing or invalid")
    if disposition in {"FROZEN", "DISCARDED"} and exit_date is None:
        warnings.append("Disposition date is unknown; zero flags cannot be classified as death")
    if v1 and source_status_folded not in {"alive", "dead", "frozen", "discarded"}:
        warnings.append("Source status needs admin review")
    fish_code = _text(cells.get("K")) if v1 else ""
    if not fish_code and dob:
        fish_code = f"V2-{dob.isoformat()}-{_text(cells.get('A'))}"
    return {
        "entity": "fish", "sourceVersion": "V1" if v1 else "V2",
        "legacyNumber": cells.get("A"), "fishCode": fish_code or None,
        "dob": dob.isoformat() if dob else None,
        "sourceStatus": source_status or None, "status": status,
        "lifeState": life_state, "disposition": disposition,
        "exitDate": exit_date.isoformat() if exit_date else None,
        "biologicalCondition": _condition(cells.get("N"), cells.get("O")) if v1 else
            ("ABNORMAL" if _text(cells.get("K")).casefold() == "abnormal" else
             "NORMAL" if _text(cells.get("K")).casefold() == "normal" else "UNDETERMINED"),
        "donorSource": cells.get("I" if v1 else "D"),
        "recipientSource": cells.get("G" if v1 else "C"),
        "sourceSpecimenCodes": [cells.get(column) for column in ("B", "C", "D", "E") if _text(cells.get(column))] if v1 else [],
        "observationCount": len(observations),
        "lastRecordedOn": observations[-1]["observedOn"] if observations else None,
        "observations": observations,
        "warnings": warnings,
    }


SPECIMEN_TYPES = {
    "whole embryo": "WHOLE_EMBRYO",
    "caudal fin clip": "CAUDAL_FIN_CLIP",
    "anal fin clip": "ANAL_FIN_CLIP",
    "left over cells": "LEFTOVER_CELLS",
    "whole adult": "WHOLE_ADULT",
}

V2_EMBRYO_STAGES = (
    "1-cell", "2-cell", "4-cell", "8-cell", "16-cell", "32-cell", "64-cell",
    "128-cell", "256-cell", "512-cell", "1k-cell", "High", "Oblong",
    "Sphere", "Dome", "30% epiboly", "50% epiboly", "Germ ring", "Shield",
    "75% epiboly", "90% epiboly",
)


def _embryo(block: dict[str, Any]) -> dict[str, Any]:
    result = block.get("result") or {}
    observed = block.get("observationTime") or {}
    degenerated = block.get("degenerated") or {}
    dead = block.get("observedDead") or {}
    condition = block.get("condition") or {}
    context = block.get("sourceContext") or {}

    def contextual(column: str) -> Any:
        return context.get(column, {}).get("value")

    experiment_date = _day(contextual("B"))
    activation_clock = _clock(contextual("AF"))
    recipient = _text(contextual("C"))
    stage_observations: list[dict[str, Any]] = []
    warnings: list[str] = []
    for index, label in enumerate(V2_EMBRYO_STAGES, 38):
        column = get_column_letter(index)
        raw = result.get(column)
        if raw is None:
            continue
        if _text(raw) in {"1", "1.0"}:
            outcome = "ALIVE"
        elif _text(raw) in {"0", "0.0"}:
            outcome = "DEGENERATED" if _text(degenerated.get(column)) else "DEAD" if _text(dead.get(column)) else "NOT_SURVIVING"
        else:
            outcome = "UNRESOLVED"
            warnings.append(f"{column}: stage value is not 1 or 0")
        observation_clock = _clock(observed.get(column))
        stage_observations.append({
            "stageLabel": label, "sourceColumn": column, "sourceValue": raw,
            "outcome": outcome, "condition": "ABNORMAL" if _text(condition.get(column)).casefold() in {"ab", "abnormal"} else "UNDETERMINED",
            "observedOn": experiment_date.isoformat() if experiment_date else None,
            "observedLocalTime": observation_clock,
            "timePrecision": "exact" if experiment_date and observation_clock else "date" if experiment_date else "unknown",
        })
    daily_survival, flag_warnings = _flags(result, 60, 209, experiment_date, "UNKNOWN", None)
    warnings.extend(flag_warnings)
    if experiment_date is None:
        warnings.append("Experiment date is missing or invalid")
    if activation_clock is None:
        warnings.append("Activation time is missing or invalid; timing metrics are unavailable")
    activation_flag = _text(result.get("AK")).casefold()
    if activation_flag not in {"1", "1.0", "0", "0.0", "y", "n", "yes", "no", ""}:
        warnings.append("Activation flag needs admin review")
    return {
        "entity": "embryo", "sourceVersion": "V2", "sourceRunningNumber": result.get("AJ"),
        "experimentDate": experiment_date.isoformat() if experiment_date else None,
        "activationLocalTime": activation_clock,
        "activationSource": result.get("AK"),
        "siteSuggestion": "MSU" if "msu" in recipient.casefold() else "KU" if recipient else None,
        "recipientSource": contextual("C"), "csofSource": contextual("V"),
        "eggCodeSource": contextual("W"), "groupSource": contextual("X"),
        "injectionSource": contextual("Y"), "lotNoSource": contextual("AC"),
        "stageObservations": stage_observations,
        "dailySurvival": daily_survival,
        "warnings": warnings,
    }


def _specimen(cells: dict[str, Any]) -> dict[str, Any]:
    code = _text(cells.get("B"))
    kind_match = re.match(r"^(CLA|CL|RT|DC)", code, re.IGNORECASE)
    material = _text(cells.get("C"))
    material_type = SPECIMEN_TYPES.get(material.casefold(), "UNKNOWN")
    return {
        "entity": "specimen", "specimenCode": code,
        "specimenKind": kind_match.group(1).upper() if kind_match else None,
        "specimenType": material_type, "sourceMaterial": material or None,
        "warnings": [] if material_type != "UNKNOWN" else ["Specimen material needs admin mapping"],
    }


V1_COUNT_STAGES = (
    "2-cell", "4-cell", "8-cell", "16-cell", "32-cell", "64-cell", "128-cell",
    "256-cell", "512-cell", "1k-cell", "High", "Oblong", "Sphere", "Dome",
    "30% epiboly", "50% epiboly", "Germ ring", "Shield", "75% epiboly",
    "90% epiboly", "Day1", "Day2",
)
MSU_COUNT_STAGES = (
    "4-cell", "8-cell", "16-cell", "32-cell", "64-cell", "256-cell",
    "512-cell", "1k-cell", "High", "Oblong", "Sphere", "Dome",
    "30% epiboly", "50% epiboly", "Germ ring", "Shield",
    "75% epiboly", "90% epiboly",
)
QC_COUNT_STAGES = (
    "2-cell", "4-cell", "8-cell", "16-cell", "32-cell", "64-cell",
    "128-cell", "256-cell", "512-cell", "1k-cell", "High", "Oblong",
    "Sphere", "Dome", "30% epiboly", "50% epiboly", "Germ ring",
    "Shield", "75% epiboly", "90% epiboly", "Bud",
)


def _aggregate(kind: str, working: dict[str, Any]) -> dict[str, Any]:
    cells = working.get("cells", working)
    context = working.get("context", {})
    observed = _day(context.get("date") if kind != "scnt_aggregate" else cells.get("A"))
    counts: list[dict[str, Any]] = []
    warnings: list[str] = []

    def add(column: str, stage: str, raw: Any, measure: str = "nAlive") -> None:
        if raw is None or _text(raw) == "":
            return
        value = _text(raw)
        if kind == "control_aggregate" and value.casefold().startswith("obs time"):
            return
        if len(value) <= 20 and re.fullmatch(r"\d+(?:\.0+)?", value):
            count = int(value.split(".", 1)[0])
            if count > 2_147_483_647:
                warnings.append(f"{column}: count exceeds the database integer limit")
                return
            counts.append({"sourceColumn": column, "stageLabel": stage, measure: count,
                           "sourceValue": raw, "observedOn": observed.isoformat() if observed else None})
        else:
            warnings.append(f"{column}: count is not a nonnegative integer ({value[:80]})")

    if kind == "legacy_lot":
        add("N", "Injected", cells.get("N"), "nTotal")
        add("Q", "Activated", cells.get("Q"), "nTotal")
        for index, stage in enumerate(V1_COUNT_STAGES, 19):
            add(get_column_letter(index), stage, cells.get(get_column_letter(index)))
        for column, stage in (("AV", "N_TU"), ("AX", "N_NHGRI"), ("AZ", "N_AB")):
            add(column, stage, cells.get(column), "nTotal")
    elif kind == "scnt_aggregate":
        add("B", "Activated", cells.get("B"), "nTotal")
        add("C", "2-cell", cells.get("C"), "nNormal")
        add("D", "2-cell", cells.get("D"), "nAbnormal")
        for index, stage in enumerate(MSU_COUNT_STAGES, 5):
            add(get_column_letter(index), stage, cells.get(get_column_letter(index)))
        add("W", "Day1", cells.get("W"), "nNormal")
        add("X", "Day1", cells.get("X"), "nAbnormal")
        for column, stage in (("Y", "Day3"), ("Z", "Swimming larva"), ("AA", "Stage indeterminate")):
            add(column, stage, cells.get(column))
    else:
        add("D", "Total", cells.get("D"), "nTotal")
        for index, stage in enumerate(QC_COUNT_STAGES, 5):
            column = get_column_letter(index)
            add(column, stage, cells.get(column))
        for day, start in ((1, 26), (2, 29), (3, 32)):
            add(get_column_letter(start + 1), f"Day{day}", cells.get(get_column_letter(start + 1)), "nNormal")
            add(get_column_letter(start + 2), f"Day{day}", cells.get(get_column_letter(start + 2)), "nAbnormal")
        for continuation in working.get("continuation", []):
            for column, raw in continuation.get("cells", {}).items():
                if len(_text(raw)) > 40 or not re.fullmatch(r"\d+\s*/\s*\d+", _text(raw)):
                    continue
                numerator, denominator = (int(part.strip()) for part in _text(raw).split("/"))
                if max(numerator, denominator) > 2_147_483_647:
                    warnings.append(f"{column}: fraction exceeds the database integer limit")
                    continue
                stage = next((label for index, label in enumerate(QC_COUNT_STAGES, 5)
                              if get_column_letter(index) == column), f"QC {column}")
                counts.append({"sourceColumn": column, "sourceRow": continuation.get("rowNo"),
                               "stageLabel": stage, "numerator": numerator,
                               "denominator": denominator, "sourceValue": raw,
                               "observedOn": observed.isoformat() if observed else None})
    total_column = "N" if kind == "legacy_lot" else "B" if kind == "scnt_aggregate" else "D"
    total = next((entry["nTotal"] for entry in counts
                  if entry["sourceColumn"] == total_column and "nTotal" in entry), None)
    if total is not None:
        for entry in counts:
            measured = entry.get("nAlive", entry.get("nNormal", entry.get("nAbnormal")))
            if measured is not None and measured > total:
                warnings.append(f"{entry['sourceColumn']}: stage count {measured} exceeds total {total}")
            if entry.get("numerator", 0) > entry.get("denominator", 0):
                warnings.append(f"{entry['sourceColumn']}: fraction numerator exceeds denominator")
    if observed is None:
        warnings.append("Observation date is missing or invalid; counts retain unknown date")
    return {"entity": "historical_stage_counts", "sourceKind": kind,
            "observedOn": observed.isoformat() if observed else None,
            "counts": counts, "warnings": warnings}


def interpret(record_kind: str, working: dict[str, Any]) -> dict[str, Any] | None:
    if record_kind == "fish":
        return _fish(working)
    if record_kind == "specimen":
        return _specimen(working)
    if record_kind == "embryo_candidate":
        return _embryo(working)
    if record_kind in {"legacy_lot", "scnt_aggregate", "control_aggregate"}:
        return _aggregate(record_kind, working)
    return None

"""Identify source cells retained for an explicit future field mapping."""

from __future__ import annotations

from typing import Any

from openpyxl.utils import get_column_letter

from .import_interpret import interpret


def deferred_cells(kind: str, source: dict[str, Any]) -> list[tuple[str, Any]]:
    if kind == "fish":
        meaning = interpret(kind, source) or {}
        if meaning.get("sourceVersion") == "V1":
            consumed = {"B", "C", "D", "E", "F", "I", "K", "L", "N", "O", "OD"}
            consumed.update(get_column_letter(index) for index in range(19, 389))
        else:
            consumed = {"B", "D", "J", "K"}
            consumed.update(get_column_letter(index) for index in range(14, 206))
        return [(column, value) for column, value in source.items() if column not in consumed]
    if kind == "specimen":
        return [(column, value) for column, value in source.items() if column not in {"B", "C"}]
    if kind in {"legacy_lot", "scnt_aggregate", "control_aggregate"}:
        meaning = interpret(kind, source) or {}
        consumed = {entry["sourceColumn"] for entry in meaning.get("counts", [])}
        cells = source.get("cells", source)
        consumed.add("A" if kind in {"scnt_aggregate", "control_aggregate"} else "B")
        result = [(f"cells.{column}", value) for column, value in cells.items() if column not in consumed]
        for entry in source.get("continuation", []):
            row = entry.get("rowNo")
            for column, value in entry.get("cells", {}).items():
                if column not in consumed:
                    result.append((f"continuation.{row}.{column}", value))
        context = source.get("context", {})
        return result + [(f"context.{key}", value) for key, value in context.items() if key != "date"]
    if kind == "sheet_metadata":
        return list(source.items())
    if kind == "embryo_candidate":
        represented_context = {
            "B",
            "C",
            "V",
            "W",
            "X",
            "Y",
            "AC",
            "AF",
            "D",
            "E",
            "F",
            "G",
            "H",
            "I",
            "J",
            "K",
            "L",
            "M",
            "N",
            "O",
            "P",
            "Q",
            "R",
            "S",
            "T",
            "U",
        }
        context = source.get("sourceContext", {})
        return [
            (f"sourceContext.{column}", entry.get("value"))
            for column, entry in context.items()
            if column not in represented_context and entry.get("value") is not None
        ]
    return []

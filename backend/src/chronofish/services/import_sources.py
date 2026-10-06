"""Read customer workbooks into immutable, source-addressed import draft rows.

This module does not guess scientific meaning. Later normalization operates on
the working copy of these rows; the source cells remain available for review.
"""

from __future__ import annotations

import csv
import io
import json
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Literal

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

MAX_FILES = 60
MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_EXPANDED_BYTES = 120 * 1024 * 1024
MAX_ROWS_PER_SHEET = 20_000
MAX_COLUMNS_PER_SHEET = 500
MAX_RECORDS_PER_JOB = 100_000
MAX_CELL_CHARS = 20_000


class SourceFormatError(ValueError):
    pass


@dataclass(slots=True)
class SourceIssue:
    severity: Literal["warning", "overridable", "blocking"]
    code: str
    message: str
    sheet_name: str
    row_no: int | None = None
    source_column: str | None = None
    source_value: str | None = None
    source_locator: str | None = None


@dataclass(slots=True)
class SourceRecord:
    sheet_name: str
    source_locator: str
    row_no: int
    record_kind: str
    source: dict[str, Any]
    working: dict[str, Any]
    natural_key: str | None = None


@dataclass(slots=True)
class ParsedSheet:
    name: str
    kind: str
    records: list[SourceRecord] = field(default_factory=list)
    issues: list[SourceIssue] = field(default_factory=list)
    row_count: int = 0


def _json_value(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, (int, float, bool, str)) or value is None:
        return value
    return str(value)


def _row_dict(row: tuple[Any, ...] | list[Any]) -> dict[str, Any]:
    return {
        get_column_letter(index): _json_value(value)
        for index, value in enumerate(row, 1)
        if value is not None and value != ""
    }


def _cell(row: dict[str, Any], column: str) -> Any:
    return row.get(column)


def _text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _identifier(value: Any) -> str | None:
    text = _text(value).casefold()
    return text or None


def _sheet_kind(name: str, headers: list[dict[str, Any]]) -> str:
    lower = name.casefold().strip()
    if lower in {"summary", "ch1", "ch2", "v1clean", "clone_small molecule"}:
        return "reconciliation"
    if lower.startswith("template") or lower.startswith("templat"):
        return "template"
    if lower.startswith("v1-msu-"):
        return "msu_aggregate"
    if lower == "qc_nbd_ivf":
        return "qc"
    if lower == "v1specimen":
        return "specimen"
    if lower == "v1fish":
        return "v1_fish"
    if lower == "v2fish":
        return "v2_fish"
    if lower == "v1raw":
        return "v1_raw"
    if lower == "v2raw" or re.match(r"^\d+\)\s*\d{4}-\d{2}-\d{2}", lower):
        return "v2_raw"
    header = " ".join(_text(value).casefold() for row in headers[:2] for value in row.values())
    if "activated time" in header and "observation time" in header:
        return "v2_raw"
    if "number of clone" in header and "day 1" in header:
        return "v2_fish"
    if "dob of clone" in header and "d1" in header:
        return "v1_fish"
    if "specimen type" in header and "code" in header:
        return "specimen"
    return "unknown"


def _record(sheet: str, row_no: int, kind: str, source: dict[str, Any], key: Any = None) -> SourceRecord:
    locator = f"A{row_no}"
    return SourceRecord(sheet, locator, row_no, kind, source, dict(source), _identifier(key))


def _extract_records(sheet: ParsedSheet, rows: list[dict[str, Any]]) -> None:
    kind = sheet.kind
    if kind == "template":
        return
    if kind == "unknown":
        sheet.issues.append(
            SourceIssue("blocking", "unknown_sheet", "Choose a supported sheet type or ignore this sheet", sheet.name)
        )
        return
    if kind == "v2_raw":
        # Each embryo is a five-row block: result, observation time,
        # degenerated, observed dead, and normal/abnormal annotations.
        marker_columns: set[str] = set()
        for row in rows[:8]:
            for column, value in row.items():
                if _text(value).casefold() == "observation time":
                    marker_columns.add(column)
        for index in range(2, len(rows) - 1):
            row = rows[index]
            following = rows[index + 1]
            has_observation_row = any(
                _text(following.get(column)).casefold() == "observation time" for column in marker_columns
            )
            if not has_observation_row:
                continue
            degen_index = index + 2
            if degen_index >= len(rows) or not any(
                _text(value).casefold() == "degenerated" for value in rows[degen_index].values()
            ):
                continue
            block = {
                "result": row,
                "observationTime": following,
                "degenerated": rows[degen_index],
                "observedDead": rows[degen_index + 1] if degen_index + 1 < len(rows) else {},
                "condition": rows[degen_index + 2] if degen_index + 2 < len(rows) else {},
            }
            row_no = index + 1
            record = _record(sheet.name, row_no, "embryo_candidate", block)
            sheet.records.append(record)
            for stage_index in range(38, 59):
                column = get_column_letter(stage_index)
                value = row.get(column)
                if value is not None and _text(value) not in {"0", "1", "0.0", "1.0"}:
                    sheet.issues.append(SourceIssue(
                        "overridable", "nonbinary_embryo_stage",
                        "Expected 1 (alive) or 0 (not surviving); review this source cell",
                        sheet.name, row_no, column, _text(value)[:1000], f"{column}{row_no}",
                    ))
        # Header-level date, control counts, and lot metadata must survive too.
        for index, row in enumerate(rows[:8], 1):
            if any(_text(value).casefold() in {"observation time", "degenerated", "observed dead"} for value in row.values()):
                continue
            if index < len(rows) and any(
                _text(value).casefold() == "observation time" for value in rows[index].values()
            ):
                continue
            if len(row) >= 5:
                sheet.records.append(_record(sheet.name, index, "sheet_metadata", row))
    elif kind == "v1_raw":
        context: dict[str, Any] = {}
        for index, row in enumerate(rows[2:], 3):
            if "B" in row:
                context["date"] = row["B"]
            if "C" in row:
                context["recipientEgg"] = row["C"]
            if "F" in row:
                context["csofLot"] = row["F"]
            if "M" not in row or "N" not in row:
                continue
            sheet.records.append(
                _record(sheet.name, index, "legacy_lot", {"cells": row, "context": dict(context)}, row.get("G"))
            )
    elif kind in {"v1_fish", "v2_fish"}:
        stage_columns = {
            column for column, title in rows[0].items()
            if re.match(r"^(?:d|day\s*)\d+\b", _text(title), re.IGNORECASE)
        } if rows else set()
        for index, row in enumerate(rows[1:], 2):
            if not _text(row.get("A")).isdigit():
                continue
            key = row.get("K") if kind == "v1_fish" else row.get("A")
            sheet.records.append(_record(sheet.name, index, "fish", row, key))
            dob_column = "F" if kind == "v1_fish" else "B"
            if not _text(row.get(dob_column)):
                sheet.issues.append(SourceIssue(
                    "blocking", "fish_dob_missing", "A fish date of birth is required",
                    sheet.name, index, dob_column, None, f"{dob_column}{index}",
                ))
            for column in stage_columns & row.keys():
                value = row[column]
                if isinstance(value, bool) or _text(value) not in {"0", "1", "0.0", "1.0"}:
                    sheet.issues.append(SourceIssue(
                        "blocking", "invalid_survival_flag", "Survival must be 1 (alive) or 0 (not surviving)",
                        sheet.name, index, column, _text(value)[:1000], f"{column}{index}",
                    ))
    elif kind == "specimen":
        for index, row in enumerate(rows[1:], 2):
            code = _text(row.get("B"))
            if re.fullmatch(r"(?:CLA|CL|RT|DC)\d+", code, re.IGNORECASE):
                sheet.records.append(_record(sheet.name, index, "specimen", row, code))
    elif kind == "qc":
        context: dict[str, Any] = {}
        for index, row in enumerate(rows[2:], 3):
            if "A" in row:
                context["date"] = row["A"]
            code = _text(row.get("B"))
            if code:
                continuation: list[dict[str, Any]] = []
                for next_index in range(index, len(rows)):
                    following = rows[next_index]
                    if "B" in following or "A" in following:
                        break
                    if following:
                        continuation.append({"rowNo": next_index + 1, "cells": following})
                sheet.records.append(
                    _record(sheet.name, index, "control_aggregate",
                            {"cells": row, "context": dict(context), "continuation": continuation}, code)
                )
    elif kind == "msu_aggregate":
        for index, row in enumerate(rows[2:], 3):
            if row.get("A") is not None and row.get("B") is not None:
                sheet.records.append(_record(sheet.name, index, "scnt_aggregate", row, row.get("A")))
    elif kind == "reconciliation":
        for index, row in enumerate(rows[1:], 2):
            if row:
                sheet.records.append(_record(sheet.name, index, "reconciliation", row))
    sheet.row_count = len(sheet.records)


def _parse_matrix(name: str, matrix: list[list[Any] | tuple[Any, ...]], formula_cells: set[str] | None = None) -> ParsedSheet:
    if len(matrix) > MAX_ROWS_PER_SHEET:
        raise SourceFormatError(f"{name}: too many rows")
    if any(len(row) > MAX_COLUMNS_PER_SHEET for row in matrix):
        raise SourceFormatError(f"{name}: too many columns")
    if any(isinstance(value, str) and len(value) > MAX_CELL_CHARS for row in matrix for value in row):
        raise SourceFormatError(f"{name}: a cell is too long")
    rows = [_row_dict(row) for row in matrix]
    sheet = ParsedSheet(name=name, kind=_sheet_kind(name, rows[:3]))
    _extract_records(sheet, rows)
    for coordinate in sorted(formula_cells or ()):
        match = re.fullmatch(r"([A-Z]+)(\d+)", coordinate)
        if not match:
            continue
        column, row_number = match.groups()
        if int(row_number) > len(rows) or column not in rows[int(row_number) - 1]:
            sheet.issues.append(
                SourceIssue(
                    "warning" if sheet.kind == "reconciliation" else "blocking",
                    "formula_cache_missing", "Excel did not save a displayed result for this formula",
                    name, int(row_number), column, None, coordinate,
                )
            )
    return sheet


def _check_xlsx_archive(content: bytes) -> None:
    if len(content) > MAX_FILE_BYTES:
        raise SourceFormatError("The workbook is too large")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 5_000 or sum(entry.file_size for entry in entries) > MAX_EXPANDED_BYTES:
                raise SourceFormatError("The workbook expands beyond the import limit")
            if not any(entry.filename == "xl/workbook.xml" for entry in entries):
                raise SourceFormatError("This is not an XLSX workbook")
    except zipfile.BadZipFile as error:
        raise SourceFormatError("This is not a valid XLSX workbook") from error


def xlsx_sheet_names(content: bytes) -> list[str]:
    _check_xlsx_archive(content)
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def parse_xlsx(content: bytes, selected_sheets: list[str]) -> list[ParsedSheet]:
    _check_xlsx_archive(content)
    if not selected_sheets or len(selected_sheets) > MAX_FILES or len(selected_sheets) != len(set(selected_sheets)):
        raise SourceFormatError("Select at least one distinct sheet")
    values = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
    formulas = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    try:
        missing = set(selected_sheets) - set(values.sheetnames)
        if missing:
            raise SourceFormatError(f"Selected sheets do not exist: {', '.join(sorted(missing))}")
        result: list[ParsedSheet] = []
        for name in selected_sheets:
            value_sheet = values[name]
            formula_sheet = formulas[name]
            formula_cells: set[str] = set()
            for row in formula_sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        formula_cells.add(cell.coordinate)
            matrix = [row for row in value_sheet.iter_rows(values_only=True)]
            result.append(_parse_matrix(name, matrix, formula_cells))
        if sum(len(sheet.records) for sheet in result) > MAX_RECORDS_PER_JOB:
            raise SourceFormatError("The selected sheets contain too many records")
        return result
    finally:
        values.close()
        formulas.close()


def decode_csv(content: bytes, encoding: str | None = None) -> tuple[str, str]:
    if len(content) > MAX_FILE_BYTES:
        raise SourceFormatError("The CSV file is too large")
    encodings = [encoding] if encoding else ["utf-8-sig", "utf-16", "cp874"]
    for candidate in encodings:
        if candidate not in {"utf-8", "utf-8-sig", "utf-16", "cp874"}:
            raise SourceFormatError("Unsupported CSV encoding")
        try:
            decoded = content.decode(candidate)
        except UnicodeDecodeError:
            continue
        if "\x00" in decoded:
            continue
        return decoded, candidate
    raise SourceFormatError("CSV text could not be decoded; select another encoding")


def csv_delimiter(decoded: str, choice: str | None = None) -> str:
    if choice is not None:
        if choice not in {",", ";", "\t", "|"}:
            raise SourceFormatError("Unsupported CSV delimiter")
        return choice
    try:
        return csv.Sniffer().sniff(decoded[:8192], delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def parse_csv(
    content: bytes, sheet_name: str, encoding: str | None = None, delimiter: str | None = None
) -> tuple[ParsedSheet, str, str]:
    decoded, used_encoding = decode_csv(content, encoding)
    used_delimiter = csv_delimiter(decoded, delimiter)
    try:
        rows = list(csv.reader(io.StringIO(decoded, newline=""), delimiter=used_delimiter))
    except csv.Error as error:
        raise SourceFormatError("CSV rows could not be read") from error
    return _parse_matrix(sheet_name, rows), used_encoding, used_delimiter


def record_json(record: SourceRecord) -> tuple[str, str]:
    return (
        json.dumps(record.source, ensure_ascii=False, separators=(",", ":")),
        json.dumps(record.working, ensure_ascii=False, separators=(",", ":")),
    )
